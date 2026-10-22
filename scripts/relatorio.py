#!/usr/bin/env python3
"""Monta o relatorio em .docx e em .pdf, com os graficos de graficos.py."""

import csv
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from fpdf import FPDF

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
GRAFICOS = os.path.join(RAIZ, "graficos")
FONTES = "/System/Library/Fonts/Supplemental"
SAIDA_DOCX = os.path.join(RAIZ, "Relatorio_Testes_de_Carga_WordPress_Locust.docx")
SAIDA_PDF = os.path.join(RAIZ, "Relatorio_Testes_de_Carga_WordPress_Locust.pdf")
MAC = "MacBookAirs-MacBook-Air"
WIN = "DESKTOP-5LB6RUD"
CENARIOS = {
    1: "Imagem de aproximadamente 1 MB",
    2: "Texto de aproximadamente 400 KB",
    3: "Imagem de aproximadamente 300 KB",
}
METRICAS = [
    ("tempo_medio", "Tempo médio de resposta"),
    ("tempo_p95", "Percentil 95 do tempo de resposta"),
    ("rps", "Requisições por segundo"),
    ("erro", "Taxa de erro"),
]


def ler_resumo(maquina):
    caminho = os.path.join(RAIZ, "locust", "resultados", maquina, "resumo.csv")
    with open(caminho, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


class DocxAlvo:
    def __init__(self):
        self.doc = Document()
        estilo = self.doc.styles["Normal"]
        estilo.font.name = "Times New Roman"
        estilo.font.size = Pt(12)

    def titulo(self, texto, subtitulo):
        paragrafo = self.doc.add_paragraph()
        paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragrafo.add_run(texto)
        run.bold = True
        run.font.size = Pt(16)
        run.font.name = "Times New Roman"
        sub = self.doc.add_paragraph()
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = sub.add_run(subtitulo)
        run.font.size = Pt(12)
        run.font.name = "Times New Roman"

    def heading(self, texto):
        self.doc.add_heading(texto, level=1)

    def paragrafo(self, texto):
        paragrafo = self.doc.add_paragraph(texto)
        paragrafo.paragraph_format.space_after = Pt(8)
        for run in paragrafo.runs:
            run.font.size = Pt(12)
            run.font.name = "Times New Roman"

    def tabela(self, cabecalho, linhas):
        table = self.doc.add_table(rows=1, cols=len(cabecalho))
        table.style = "Table Grid"
        for i, nome in enumerate(cabecalho):
            table.rows[0].cells[i].text = nome
        for linha in linhas:
            cells = table.add_row().cells
            for i, valor in enumerate(linha):
                cells[i].text = str(valor)
        self.doc.add_paragraph()

    def figura(self, caminho, legenda):
        if not os.path.isfile(caminho):
            return
        self.doc.add_picture(caminho, width=Cm(15.5))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = self.doc.add_paragraph(legenda)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cap.runs:
            run.italic = True
            run.font.size = Pt(10)
            run.font.name = "Times New Roman"

    def salvar(self, caminho):
        self.doc.save(caminho)


class PdfAlvo(FPDF):
    def __init__(self):
        super().__init__(format="A4")
        self.set_auto_page_break(auto=True, margin=18)
        self.add_font("TNR", "", os.path.join(FONTES, "Times New Roman.ttf"))
        self.add_font("TNR", "B", os.path.join(FONTES, "Times New Roman Bold.ttf"))
        self.add_font("TNR", "I", os.path.join(FONTES, "Times New Roman Italic.ttf"))
        self.add_page()
        self.set_font("TNR", "", 12)

    def footer(self):
        self.set_y(-12)
        self.set_font("TNR", "", 9)
        self.cell(0, 8, str(self.page_no()), align="C")

    def titulo(self, texto, subtitulo):
        self.set_font("TNR", "B", 16)
        self.multi_cell(0, 8, texto, align="C")
        self.ln(2)
        self.set_font("TNR", "", 12)
        self.multi_cell(0, 6, subtitulo, align="C")
        self.ln(4)

    def heading(self, texto):
        self.ln(2)
        self.set_font("TNR", "B", 14)
        self.multi_cell(0, 8, texto)
        self.ln(1)
        self.set_font("TNR", "", 12)

    def paragrafo(self, texto):
        self.set_font("TNR", "", 12)
        self.multi_cell(0, 6, texto)
        self.ln(2)

    def tabela(self, cabecalho, linhas):
        dados = [cabecalho, *[[str(v) for v in linha] for linha in linhas]]
        self.set_font("TNR", "", 8)
        with self.table(width=190, col_widths=tuple(1 for _ in cabecalho), line_height=5) as table:
            for i, linha in enumerate(dados):
                row = table.row()
                for valor in linha:
                    row.cell(valor, align="C" if i else "C")
        self.ln(3)
        self.set_font("TNR", "", 12)

    def figura(self, caminho, legenda):
        if not os.path.isfile(caminho):
            return
        if self.get_y() > 180:
            self.add_page()
        self.image(caminho, w=180)
        self.ln(1)
        self.set_font("TNR", "I", 10)
        self.multi_cell(0, 5, legenda, align="C")
        self.ln(3)
        self.set_font("TNR", "", 12)

    def salvar(self, caminho):
        self.output(caminho)


def figura_path(maquina, cenario, sufixo):
    return os.path.join(GRAFICOS, maquina, f"cenario{cenario}_{sufixo}.png")


def escrever(alvo, mac, win):
    alvo.titulo(
        "Testes de carga de um cluster WordPress com Locust",
        "Computação Distribuída\nThiago e Jean",
    )
    alvo.heading("1. Objetivo")
    alvo.paragrafo(
        "O trabalho mede o desempenho de um cluster WordPress quando variam o número de instâncias da aplicação, a quantidade de usuários simultâneos e o tipo de conteúdo acessado. A carga é gerada pelo Locust, executado na mesma máquina e na mesma rede Docker da aplicação, como pedido pelo professor. O resultado de cada cenário é apresentado em gráficos de barras: o eixo X é o número de instâncias do WordPress e o eixo Y é a métrica coletada, com uma barra para cada nível de carga."
    )
    alvo.heading("2. Arquitetura")
    alvo.paragrafo(
        "O ambiente replica o cluster do Trabalho 2. Um contêiner MySQL 5.7 guarda o banco. De uma a três instâncias wordpress:5.4.2-php7.2-apache atendem o site e compartilham o diretório de arquivos. O Nginx 1.19.0 publica a porta 80 e distribui as requisições entre as instâncias ativas, conforme o arquivo nginx/nginx-N.conf. O Locust 1.2.3, imagem grubykarol/locust, entra no mesmo docker-compose e aponta para http://nginx. O gerador de carga, portanto, disputa CPU e memória com o próprio sistema medido."
    )
    alvo.paragrafo(
        "Cada usuário virtual espera entre 1 e 2 segundos, solicita o post do cenário e baixa as imagens de wp-content/uploads referenciadas na página, no máximo três. As estatísticas do Locust contam requisições HTTP, incluindo a página e as imagens, e não apenas visitas completas. O cabeçalho Host é localhost, e redirecionamentos absolutos são convertidos em caminho relativo para continuarem passando pelo Nginx."
    )
    alvo.heading("3. Cenários e cargas")
    alvo.paragrafo("Os três cenários do enunciado foram publicados como posts públicos, com os mesmos identificadores usados na outra máquina.")
    alvo.tabela(["Cenário", "Post", "Conteúdo medido"], [
        ["1", "5", "Imagem de 1.084.779 bytes (cerca de 1 MB)"],
        ["2", "8", "Texto de 409.600 bytes (400 KB)"],
        ["3", "10", "Imagem de 306.626 bytes (cerca de 300 KB)"],
    ])
    alvo.paragrafo(
        "A bateria comparativa usa os mesmos três níveis de usuários nas duas máquinas e em todos os valores de N: 32 (mínima), 80 (média) e 160 (máxima). Esses valores saíram da calibração de N = 3 registrada em DESKTOP-5LB6RUD e foram reaplicados aqui para que a comparação não misture cargas diferentes. Cada execução dura 60 segundos, com pausa de 15 segundos entre elas. A carga mínima e a média são válidas com 0% de erro. A máxima é válida com até 10% de erro."
    )
    alvo.paragrafo("Foram medidas três quantidades de instâncias, N = 1, 2 e 3. A combinação produz 27 execuções por máquina: 3 cenários × 3 cargas × 3 arquiteturas.")
    alvo.heading("4. Máquinas")
    alvo.tabela(["Máquina", "Processador", "Memória", "Sistema", "Papel no teste"], [
        ["MacBookAirs-MacBook-Air", "Apple M1 (8 núcleos)", "8 GB; cerca de 4 GB no Docker", "macOS, imagens amd64 emuladas", "27 execuções"],
        ["DESKTOP-5LB6RUD", "Intel Core i3-7020U", "8 GB", "Windows 10", "Calibração completa; bateria parcial"],
    ])
    alvo.paragrafo(
        "As imagens oficiais usadas no Trabalho 2 são amd64. No Mac elas rodam por emulação. Essa diferença de processador e o limite de memória do Docker fazem parte do resultado e impedem tratar um tempo menor como mérito exclusivo do número de instâncias."
    )
    alvo.heading("5. Resultados neste Mac")
    alvo.paragrafo("As 27 execuções foram concluídas e todas respeitaram a regra de erro: mínima e média com 0%, máxima com no máximo 4,81%. A tabela abaixo reúne tempo médio, percentil 95, requisições por segundo e erro.")
    alvo.tabela(
        ["N", "Cenário", "Carga", "Usuários", "Média (ms)", "p95 (ms)", "Req/s", "Erro (%)"],
        [[l["n"], l["cenario"], l["carga"], l["usuarios"], l["media_ms"], l["p95_ms"], l["rps"], l["erro_pct"]] for l in mac],
    )
    alvo.paragrafo("Os gráficos seguem o formato pedido: instâncias no eixo X e uma barra por carga. Há um conjunto para cada cenário e para cada métrica.")
    numero = 1
    for cenario, nome in CENARIOS.items():
        for sufixo, metrica in METRICAS:
            alvo.figura(
                figura_path(MAC, cenario, sufixo),
                f"Figura {numero}. {nome}: {metrica.lower()} neste Mac.",
            )
            numero += 1
    alvo.heading("6. O que esses gráficos mostram")
    alvo.paragrafo(
        "Na carga mínima, 32 usuários, o tempo médio fica baixo nos três cenários e quase não muda quando se acrescentam instâncias. Com uma instância, a imagem de 1 MB responde em 24 ms e a de 300 KB em 30 ms. O texto de 400 KB é mais pesado, 372 ms, porque cada resposta carrega centenas de kilobytes de HTML."
    )
    alvo.paragrafo(
        "Na carga média, 80 usuários, a segunda instância muda o resultado da imagem de 1 MB: o tempo médio cai de 422 ms para 41 ms e a vazão sobe de 84 para 143 requisições por segundo. A terceira instância não repete esse ganho (59 ms e 139 req/s). No texto de 400 KB ocorre o contrário: o tempo médio passa de 470 ms com uma instância para 622 ms com duas e 848 ms com três, e a vazão diminui. A imagem de 300 KB já era rápida com uma instância na carga média (72 ms) e permanece nessa faixa."
    )
    alvo.paragrafo(
        "Na carga máxima, 160 usuários, o serviço entra em saturação em todos os cenários. O tempo médio fica entre 2,6 s e 8,5 s, o percentil 95 chega a dezenas de segundos e a vazão fica abaixo da carga média. O erro aparece só nesse nível e permanece dentro do limite de 10%: 2,80% e 2,62% com uma instância nas imagens, 4,81% e 2,73% com duas, e 0% com três. Mais instâncias não aumentam a vazão máxima. MySQL, Nginx e Locust compartilham o mesmo computador, e o Docker desta máquina dispõe de cerca de 4 GB. A partir de certo ponto, o recurso escasso deixa de ser o número de processos PHP."
    )
    alvo.heading("7. Resultados na outra máquina")
    alvo.paragrafo(
        "DESKTOP-5LB6RUD calibrou as cargas para N = 1, 2 e 3. Com uma instância, a calibração aceitou uma máxima de 480 usuários. Com duas e com três, a máxima caiu para 160. A bateria comparativa foi então fixada em 32, 80 e 160 usuários para os três valores de N. No repositório, o resumo dessa bateria contém 5 das 27 execuções, todas com N = 1 e erro 0%: as três cargas da imagem de 1 MB e as cargas mínima e média do texto de 400 KB. Os gráficos abaixo usam somente essas linhas. Não há medição, neste arquivo, para a imagem de 300 KB nem para N = 2 e N = 3."
    )
    alvo.tabela(
        ["N", "Cenário", "Carga", "Usuários", "Média (ms)", "p95 (ms)", "Req/s", "Erro (%)"],
        [[l["n"], l["cenario"], l["carga"], l["usuarios"], l["media_ms"], l["p95_ms"], l["rps"], l["erro_pct"]] for l in win],
    )
    for cenario in sorted({int(l["cenario"]) for l in win}):
        for sufixo, metrica in METRICAS:
            alvo.figura(
                figura_path(WIN, cenario, sufixo),
                f"Figura {numero}. {CENARIOS[cenario]}: {metrica.lower()} em DESKTOP-5LB6RUD. Só há pontos para N = 1.",
            )
            numero += 1
    alvo.heading("8. Comparação entre as máquinas")
    alvo.paragrafo(
        "Onde as duas baterias se encontram, N = 1 e as mesmas quantidades de usuários, o Mac responde mais rápido e sustenta mais requisições por segundo. Na imagem de 1 MB com 32 usuários, o tempo médio é 24 ms neste Mac e 1.529 ms no desktop; a vazão é 59,6 contra 14,6 req/s. Com 160 usuários, os tempos médios são 4,7 s e 11,5 s. No texto de 400 KB com 32 usuários, são 372 ms contra 3,7 s. As duas máquinas mostram o mesmo formato de curva: a carga mínima ainda cabe, a média alonga a resposta e a máxima empurra o percentil 95 para vários segundos."
    )
    alvo.paragrafo(
        "A diferença absoluta não separa o efeito do WordPress do efeito do hardware. O desktop é um i3 de dois núcleos com 8 GB, e o Locust roda no mesmo Windows. O Mac é um M1, mas executa as imagens amd64 por emulação e limita o Docker a cerca de 4 GB. A comparação válida é a forma do resultado sob a mesma carga, não um ranking de qual computador é mais rápido."
    )
    alvo.paragrafo(
        "A calibração do desktop também mostra que aumentar instâncias pode reduzir a carga máxima aceitável: de 480 usuários com N = 1 para 160 com N = 2 e N = 3. Isso combina com o que a bateria deste Mac mediu na carga de 160 usuários e no texto grande: mais réplicas do WordPress, no mesmo host, passam a competir com o banco e com o gerador de carga."
    )
    alvo.heading("9. Conclusão")
    alvo.paragrafo(
        "Os três cenários foram executados com três cargas crescentes e três quantidades de instâncias. Neste Mac, as 27 execuções são válidas pela regra de erro do trabalho. A segunda instância melhora a imagem de 1 MB na carga média. A terceira não traz ganho claro, e no texto de 400 KB o acréscimo de instâncias piora tempo e vazão. Na carga de 160 usuários o cluster satura nas duas máquinas, com percentil 95 muito acima da média, embora o erro tenha ficado abaixo de 10%."
    )
    alvo.paragrafo(
        "O Locust no mesmo host é parte do experimento: parte da CPU e da memória medida pertence ao gerador, não ao WordPress. Por isso o aumento de instâncias não deve ser lido como aumento garantido de capacidade. A bateria do desktop ainda precisa ser completada até 27 linhas para que os gráficos de N = 2 e N = 3 daquela máquina tenham o mesmo recorte dos gráficos deste Mac."
    )


def main():
    mac = ler_resumo(MAC)
    win = ler_resumo(WIN)
    docx = DocxAlvo()
    escrever(docx, mac, win)
    docx.salvar(SAIDA_DOCX)
    pdf = PdfAlvo()
    escrever(pdf, mac, win)
    pdf.salvar(SAIDA_PDF)
    print(SAIDA_DOCX)
    print(SAIDA_PDF)


if __name__ == "__main__":
    main()
